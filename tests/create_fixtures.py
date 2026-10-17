"""Generate disposable upload fixtures, never user document copies."""
from pathlib import Path
from tempfile import gettempdir
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches
content = [('Target audience', 'A target audience shares needs and behaviors.'), ('Test budget', 'Start small and compare channel conversion rates.')]
directory = Path(gettempdir())
c = canvas.Canvas(str(directory / 'linba-fixture.pdf'), pagesize=(960, 540))
c.setAuthor('Codex')
for title, body in content:
    c.setFont('Helvetica-Bold', 30)
    c.drawString(50, 400, title)
    c.setFont('Helvetica', 20)
    c.drawString(50, 320, body)
    c.showPage()
c.save()
p = Presentation()
p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
for title, body in content:
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
p.save(directory / 'linba-fixture.pptx')
print('Upload fixtures created in the system temporary directory.')
